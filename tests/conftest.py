"""
Shared test fixtures.

IMPORTANT, and worth reading if you're reviewing this test suite: these tests
run fully offline. The real embedding model (sentence-transformers, downloaded
from huggingface.co) and the real Groq API are both replaced with small, fast,
deterministic fakes below. That means this suite verifies the *pipeline*
end-to-end (upload -> chunk -> store -> retrieve -> filter -> persist ->
delete -> conversation memory -> ...) very thoroughly, but it does NOT prove
the real MiniLM embedding model or the real Groq API behave a particular way.
Before relying on this in production, do at least one real upload + real
question with a valid GROQ_API_KEY and network access.
"""
from __future__ import annotations

import pytest
from fastapi.testclient import TestClient
from langchain_core.embeddings import Embeddings


class FakeEmbeddings(Embeddings):
    """Deterministic bag-of-words style embedding: good enough to test
    ordering/filtering/scoping logic, not a real semantic model."""

    _VOCAB = [
        "cat", "dog", "python", "invoice", "contract", "recipe", "grape",
        "warranty", "architecture", "retrieval", "revenue", "clause",
    ]

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._vec(t) for t in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._vec(text)

    def _vec(self, text: str) -> list[float]:
        low = text.lower()
        return [float(low.count(word)) for word in self._VOCAB]


class FakeLLMResponse:
    def __init__(self, content: str):
        self.content = content


class FakeChatModel:
    """Stands in for ChatGroq. `invoke` returns a canned/echo answer;
    `stream` yields it word by word so the streaming code path is exercised
    without a real network connection."""

    def __init__(self, canned_answer: str | None = None):
        self.canned_answer = canned_answer
        self.last_messages = None

    def invoke(self, messages):
        self.last_messages = messages
        return FakeLLMResponse(self._answer_for(messages))

    def stream(self, messages):
        self.last_messages = messages
        text = self._answer_for(messages)
        for word in text.split(" "):
            yield FakeLLMResponse(word + " ")

    def _answer_for(self, messages) -> str:
        if self.canned_answer is not None:
            return self.canned_answer
        # Default: a deterministic "echo" answer that proves the LLM actually
        # received the retrieved context, without needing a real model.
        last_human = messages[-1].content if messages else ""
        return f"[fake answer based on context] {last_human[:120]}"


@pytest.fixture()
def client(tmp_path, monkeypatch):
    """A fully isolated TestClient: its own temp document store, its own temp
    Chroma dir, its own temp SQLite file, fake embeddings, fake LLM."""
    from backend import config

    monkeypatch.setattr(config, "DOCUMENT_STORE_DIR", tmp_path / "document_store")
    monkeypatch.setattr(config, "CHROMA_DB_DIR", tmp_path / "chroma_db")
    monkeypatch.setattr(config, "SQLITE_DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(config, "FRONTEND_DIR", tmp_path / "no_frontend_here")
    monkeypatch.setattr(config, "GROQ_API_KEY", "test-fake-key")
    monkeypatch.setattr(config, "MAX_UPLOAD_SIZE_MB", 1)
    monkeypatch.setattr(config, "MAX_UPLOAD_SIZE_BYTES", int(1 * 1024 * 1024))

    from backend.services import embedding_service, vector_store, rag_service

    embedding_service.get_embeddings.cache_clear()
    vector_store.get_store.cache_clear()
    rag_service.get_llm.cache_clear()
    monkeypatch.setattr(embedding_service, "get_embeddings", lambda: FakeEmbeddings())
    fake_llm = FakeChatModel()
    monkeypatch.setattr(rag_service, "get_llm", lambda: fake_llm)

    from backend.main import app

    with TestClient(app) as test_client:
        yield test_client


class AlwaysFailingChatModel:
    """Every call raises -- simulates total LLM unavailability (e.g. a
    retired/misspelled model name, exactly what actually happened with the
    Groq API when this project was first run for real)."""

    def invoke(self, messages):
        raise RuntimeError("simulated Groq failure: model not found")

    def stream(self, messages):
        raise RuntimeError("simulated Groq failure: model not found")
        yield  # pragma: no cover -- unreachable; keeps this a generator function


class CondenseFailsChatModel:
    """Fails only on the question-condensation call (detected by the
    distinctive prompt marker), succeeds normally otherwise -- lets tests
    verify a condensation failure degrades gracefully instead of blocking
    the whole answer."""

    def __init__(self, canned_answer: str = "[answer despite condensation failure]"):
        self.canned_answer = canned_answer
        self.last_messages = None

    def invoke(self, messages):
        self.last_messages = messages
        last = messages[-1].content if messages else ""
        if "Standalone question:" in last:
            raise RuntimeError("simulated condensation failure")
        return FakeLLMResponse(self.canned_answer)

    def stream(self, messages):
        self.last_messages = messages
        for word in self.canned_answer.split(" "):
            yield FakeLLMResponse(word + " ")


@pytest.fixture()
def sample_txt_bytes() -> bytes:
    return (
        b"DocQuery Test Document\n\n"
        b"This document describes a three-stage retrieval architecture "
        b"for question answering systems. The architecture was proposed "
        b"by the engineering team in 2024. It combines dense retrieval, "
        b"a reranking stage, and a generation stage.\n\n"
        b"The system's monthly revenue grew after the new pricing model launched."
    )


def make_pdf_bytes(pages: list[str]) -> bytes:
    """Build a real, small, text-bearing PDF (via reportlab) for tests --
    NOT a production dependency, only used here to create realistic fixtures."""
    import io

    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    for page_text in pages:
        text_object = c.beginText(72, 720)
        for line in page_text.split("\n"):
            text_object.textLine(line)
        c.drawText(text_object)
        c.showPage()
    c.save()
    return buffer.getvalue()


def make_docx_bytes(sections: list[tuple[str, list[str]]], tables: list[list[list[str]]] | None = None) -> bytes:
    """Build a real .docx (via python-docx) with headings + paragraphs + optional tables."""
    import io

    from docx import Document as DocxFile

    doc = DocxFile()
    for heading, paragraphs in sections:
        doc.add_heading(heading, level=1)
        for para in paragraphs:
            doc.add_paragraph(para)
    for table_rows in tables or []:
        table = doc.add_table(rows=0, cols=len(table_rows[0]))
        for row_values in table_rows:
            row = table.add_row()
            for cell, value in zip(row.cells, row_values):
                cell.text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
